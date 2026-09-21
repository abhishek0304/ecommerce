from pathlib import Path
import json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from api_catalog import APIS, ORDER, OID, T
from narrative import CHAPTERS, QUESTIONS

ROOT=Path(__file__).resolve().parents[2]
BUILD=Path(__file__).parent
OUT=ROOT/'docs/Ecommerce_Application_Architecture_API_and_Interview_Guide.docx'
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.08
for n,size in [(1,19),(2,14),(3,12)]:
    st=d.styles[f'Heading {n}']; st.font.size=Pt(size); st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(6)
d.styles['Title'].font.size=Pt(32)
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
st=d.styles.add_style('Code',1); st.font.name='Consolas'; st.font.size=Pt(9)
st.paragraph_format.space_after=Pt(0); st.paragraph_format.line_spacing=1
st.paragraph_format.keep_together=True
header=s.header.paragraphs[0]; header.text='ECOMMERCE APPLICATION   |   TECHNICAL HANDBOOK'; header.style='Caption'
foot=s.footer.paragraphs[0]; foot.text='Source snapshot 15 September 2026                                             '
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); foot._p.append(f)

def p(text):
    for v in text.split('\n'):
        if v.strip(): d.add_paragraph(v.strip())
def h(text,level=1): d.add_heading(re.sub(r'[^\w ]',' ',text).replace('_',' '),level)
def label(text):
    z=d.add_paragraph(); z.paragraph_format.keep_with_next=True; z.add_run(text).bold=True
def compact(x,indent=0):
    plain=json.dumps(x,ensure_ascii=False,separators=(', ', ': '))
    if len(plain)+indent<=96: return plain
    if isinstance(x,dict): return '{\n'+',\n'.join(' '*(indent+2)+json.dumps(k)+': '+compact(v,indent+2) for k,v in x.items())+'\n'+' '*indent+'}'
    if isinstance(x,list): return '[\n'+',\n'.join(' '*(indent+2)+compact(v,indent+2) for v in x)+'\n'+' '*indent+']'
    return plain
def code(x):
    text=x if isinstance(x,str) else json.dumps(x,ensure_ascii=False,indent=2)
    # Compact individual primitive objects and arrays while retaining standard JSON.
    if not isinstance(x,str):
        text=compact(x)
    lines=text.splitlines()
    for i,line in enumerate(lines):
        z=d.add_paragraph(line,'Code')
        if len(lines)>1 and (i<2 or i==len(lines)-2):
            z.paragraph_format.keep_with_next=True
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F2F4F7'); z._p.get_or_add_pPr().append(sh)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
def table(headers,rows):
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=True
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    pr=t.rows[0]._tr.get_or_add_trPr(); pr.append(OxmlElement('w:tblHeader'))
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for i,row in enumerate(t.rows):
        for c in row.cells:
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'CBD5E1');borders.append(el)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'233D53');pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for r in para.runs:
                    r.font.size=Pt(9)
                    if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    d.add_paragraph()

def diagram():
    im=Image.new('RGB',(1500,880),'white'); a=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/calibri.ttf',28)
    def box(x,y,w,txt):
        a.rounded_rectangle((x,y,x+w,y+88),12,fill='#edf3f8',outline='#3b607b',width=3)
        a.multiline_text((x+w/2,y+44),txt,font=font,fill='#182e40',anchor='mm',align='center')
    def arrow(x,y,u,v):
        a.line((x,y,u,v),fill='#51738b',width=4)
        a.polygon([(u,v),(u-8,v-14),(u+8,v-14)],fill='#51738b')
    box(550,10,400,'Browser storefront and admin')
    arrow(750,98,750,145);box(550,145,400,'API gateway 8081')
    names=['User 8082','Product 8083','Cart 8084','Order 8085','Notification 8086']
    for i,n in enumerate(names):
        x=20+i*298; a.line((750,233,750,265),fill='#51738b',width=3);a.line((x+140,265,750,265),fill='#51738b',width=3)
        arrow(x+140,265,x+140,310);box(x,310,280,n)
        arrow(x+140,398,x+140,435);box(x,435,280,['userdb','productdb','cartdb','orderdb','notificationdb'][i])
    box(20,595,280,'Redis product cache');box(385,595,700,'Kafka user events and order events');box(1150,595,330,'Razorpay test API')
    box(260,750,450,'Eureka 8761 discovery');box(780,750,450,'Config server 8080')
    im.save(BUILD/'architecture.png'); d.add_picture(str(BUILD/'architecture.png'),width=Inches(7))
    p('Figure 1. Each business service owns its database. Order-service calls user, cart, product and Razorpay synchronously; order events reach notification-service through Kafka, while user events are published to a separate topic. Redis supports product reads. Discovery and configuration support the runtime rather than owning commerce data.')

d.add_paragraph('Ecommerce Application', 'Title')
d.add_paragraph('Architecture API Reference and Interview Guide','Subtitle')
p('Prepared for the project developer\nBased on the repository at D:\\spring projects\\ecommerce\nSource review date 15 September 2026')
p('This handbook explains how the implemented application works, how requests cross its services, and why its validation, authorization, concurrency and recovery conditions exist. It includes concrete request and response examples for all application routes in 63 reference entries covering 64 handler methods, compatibility aliases, internal contracts and 60 project-specific interview answers.')
p('The application uses local database transactions and durable recovery to coordinate commerce operations across services. Its most important interview topic is how checkout remains safe when HTTP calls, payment confirmations or event delivery fail at different times.')
p('Examples are illustrative contracts derived from source code, not captured production traffic. IDs, customer information, tokens, timestamps and provider references are fictional. Operational and security gaps are called out where the source does not support a stronger claim.')
d.add_page_break();h('Guide contents')
for text in ['Understanding the application','Architecture and service ownership','Control flow from registration to delivery','Reading states and data contracts','Complete application API reference','Infrastructure and provider interfaces','Conditions and why they exist','Operations testing and limitations','Interview questions and project answers','Source map and practical walkthrough']:
    p(text)
p('Use the Word Navigation pane to jump to any chapter, API operation or interview question. API sections identify access requirements, request bodies, success structures, important conditions, expected errors and the implementing classes. Empty bodies are stated explicitly. Framework generated error details can vary by Spring version and security filter path.')

for ci,(title,sections) in enumerate(CHAPTERS[:4]):
    d.add_page_break();h(title)
    if ci==1:
        diagram()
        table(['Service','Port','Responsibility'],[['service-registry','8761','Eureka service discovery'],['config-server','8080','Central configuration source'],['api-gateway','8081','Routing and static browser UI'],['user-service','8082','Accounts authentication addresses preferences'],['product-service','8083','Catalog stock reservations Redis cache'],['cart-service','8084','Versioned customer cart'],['order-service','8085','Checkout payments fulfillment coupons reviews wishlist returns'],['notification-service','8086','Kafka events and customer inbox']])
    for title,body in sections:h(title,2);p(body)

d.add_page_break();h('Complete application API reference')
p('Default base URL is http://localhost:8081 for explicit /api routes. Internal requests use the stated service port and X-Service-Key. Legacy aliases and /hello are service controller mappings; do not assume they are explicit public gateway routes. Authentication is enforced by each business service. JSON requests use Content-Type: application/json, and JSON responses normally use application/json. Replace concrete IDs in paths with returned values. Pagination examples show a typical Spring Page serialization; metadata ordering and framework details are not a stable custom DTO promise.')
p('Authorization: Bearer <ACCESS_TOKEN> is required unless the entry states Public, a refresh token, a webhook signature or a service key. Admin tokens require ROLE_ADMIN. Public means authentication is not required; a malformed Authorization header may still be rejected by a service filter. A response body of null in this guide is never a replacement for an explicitly stated empty response body.')
h('Common error structures',2)
p('Errors are not globally standardized. User-service custom business failures return a message object. Product-service advice returns an ApiError. Cart and order business exceptions use ProblemDetail in relevant advice paths. Security entry points, request binding failures and framework validation can use different structures. Treat the endpoint error notes as status and condition guidance rather than assuming every failure shares the following body.')
code({'message':'The request conflicts with existing data'})
code({'timestamp':T,'status':404,'error':'Not Found','message':'Product not found','path':'/api/v1/products/999','fieldErrors':{}})
code({'type':'about:blank','title':'Conflict','status':409,'detail':'Cart changed concurrently. Reload the cart and retry.','instance':'/api/cart/items'})
last=None
for i,a in enumerate(APIS,1):
    start_paragraph=len(d.paragraphs)
    if a['group']!=last:
        d.add_page_break();h(a['group'],2);last=a['group']
    h(f"API {i} {a['name']}",3)
    code(a['method']+' '+a['route'])
    p('Access: '+a['access'])
    if a['aliases']:p('Additional mappings: '+'; '.join(a['aliases']))
    if a['headers']:code('\n'.join(k+': '+v for k,v in a['headers'].items()))
    label('Request body')
    if a['request'] is None:p('No request body. Parameters, when present, are shown in the URL or headers above.')
    else:code(a['request'])
    label(f"Success response example   HTTP {a['status']}")
    if a['response'] is None:p('Empty response body.')
    else:code(a['response'])
    label('Behavior and reasons');p(a['rules'])
    label('Other responses and failure conditions');p(a['errors'])
    p('Source: '+a['source'])
    if i==48:
        for para in d.paragraphs[start_paragraph:]:
            if para.style.name=='Normal': para.paragraph_format.space_after=Pt(2)

d.add_page_break();h('Infrastructure and provider interfaces')
h('Razorpay checkout response variant',2)
p('Order responses use the same OrderView structure shown above. In a usable online checkout, paymentMethod is RAZORPAY, status is PENDING_PAYMENT, paymentStatus is UNPAID, and checkout contains the object below. The amount is an integer number of paise. The browser receives the public test key ID; the secret key stays on the server. Other order fields remain present, including expiry and refund state.')
code({'keyId':'rzp_test_EXAMPLE','razorpayOrderId':'order_demo1','amount':180000,'currency':'INR','testMode':True})
h('Outbound provider requests',2)
p('The following seven HTTP operations are emitted by RazorpayClient in this repository. This is a source contract description, not a claim to exhaustively document the current external provider API. Base URL: https://api.razorpay.com/v1. Every call uses HTTP Basic authentication with configured key ID and secret. Never embed the secret in browser JavaScript or this document. Responses below include the fields consumed by this application; Razorpay can return additional fields.')
provider_order={'id':'order_demo1','receipt':OID,'amount':180000,'currency':'INR','status':'paid'}
payment={'id':'pay_demo1','order_id':'order_demo1','amount':180000,'currency':'INR','status':'captured'}
refund={'id':'rfnd_demo1','payment_id':'pay_demo1','amount':180000,'currency':'INR','status':'processed'}
for route,req,res in [('POST /orders',{'receipt':OID,'amount':180000,'currency':'INR','partial_payment':False},dict(provider_order,status='created')),('GET /orders?receipt='+OID+'&count=100',None,{'items':[provider_order]}),('GET /orders/order_demo1',None,provider_order),('GET /orders/order_demo1/payments',None,{'items':[payment]}),('GET /payments/pay_demo1',None,payment),('POST /payments/pay_demo1/refund',{'amount':180000,'speed':'normal','receipt':OID},refund),('GET /refunds/rfnd_demo1',None,refund)]:
    label(route);p('Request body: none.' if req is None else 'Request JSON')
    if req:code(req)
    if route.startswith('POST /payments'):p('Additional header: X-Refund-Idempotency: <persisted refund attempt UUID>. The receipt example represents the attempt ID, not the order ID; these are distinct values in a real request.')
    p('Successful response consumed fields');code(res)
h('Framework and static HTTP surfaces',2)
table(['Surface','Request and response','Scope'],[['Gateway storefront','GET / and /index.html → HTML; referenced JS and CSS assets → text resources','Browser interface not JSON business APIs'],['Actuator','GET /actuator/health → 200 {"status":"UP"}; unhealthy status can be 503','Container management port 9090; exposure and health details depend on profile'],['Actuator information','GET /actuator/info → configured information object, possibly {}','Container exposes health and info; IDE gateway configuration is broader'],['Config Server','GET /order-service/default → Environment JSON','name profiles label version state and propertySources; property values omitted here to avoid publishing configuration secrets'],['Eureka registry','GET /eureka/apps with Accept application/json → applications object','Framework registration heartbeat and deregistration support discovery; not customer commerce operations'],['HTTP protocol methods','HEAD on supported GET routes omits body; OPTIONS is framework handled','CORS and security policy depend on configuration; no custom business body']])
p('Config Server can also expose the framework YAML and properties representations such as /order-service-default.yml. Eureka clients register instance metadata, renew leases and unregister through its framework protocol. These automatically supplied endpoints are distinct from the complete application controller catalog. Do not expose registry, configuration or internal inventory endpoints to customers merely because discovery can resolve their service IDs.')
h('Event request and response equivalents',2)
p('Kafka messages have no synchronous HTTP response. Producer acknowledgement, consumer acknowledgement and a persisted inbox row occur at separate times. The order outbox publishes an envelope like this, keyed by orderId. eventId deduplicates delivery and sequence identifies order-local progression.')
code({'schemaVersion':1,'eventId':'22222222-2222-4222-8222-222222222222','orderId':OID,'userId':1,'sequence':1,'type':'OrderConfirmed','occurredAt':T,'status':'CONFIRMED','paymentStatus':'UNPAID','refundStatus':'NONE','carrier':None,'trackingNumber':None})
p('Recognized order event types include OrderConfirmed, OrderFailed, PaymentPending, PaymentCaptured, PaymentCollected, OrderExpired, OrderCancelled, RefundRequested, RefundPending, RefundProcessed, RefundFailed, OrderProcessing, OrderShipped, OrderDelivered, ReturnRequested, ReturnApproved, ReturnRejected and ReturnReceived. The consumer validates schema version, UUID identifiers, positive userId and sequence. Unknown types are ignored. A duplicate eventId does not create a second inbox row.')
code({'type':'PasswordResetOtpRequested','userId':1,'email':'asha@example.com','occurredAt':T})
p('User events carry metadata and do not contain the OTP. They are published directly rather than through the order outbox. This sample therefore does not imply functioning email delivery or equivalent crash durability.')

for title,sections in CHAPTERS[4:]:
    d.add_page_break();h(title)
    for title,body in sections:h(title,2);p(body)
d.add_page_break();h('Interview questions and project answers')
p('Answers below describe the checked-in implementation. In an interview, explain the concrete mechanism, then the failure case it addresses, and finally the remaining limitation. Do not claim live payment processing, universal exactly-once delivery or complete access-token revocation.')
for i,(question,answer) in enumerate(QUESTIONS,1):h(f'Question {i} {question}',2);p(answer)
d.add_page_break();h('Source map and practical walkthrough')
p('Paths below are relative to D:\\spring projects\\ecommerce. Class names in API entries can be located under these roots. This source map makes the handbook reviewable against the implementation rather than treating examples as a separate API specification.')
for name,path in [('User controllers and account logic','user-service/user-service/src/main/java/com/ecommerce/user_service'),('Catalog cache and stock reservations','product-service/src/main/java/com/ecommerce/product_service'),('Versioned cart and checkout snapshot','cart-service/src/main/java/com/ecommerce/cart_service'),('Checkout payment and recovery coordinator','order-service/src/main/java/com/ecommerce/order_service'),('Notification ingestion and inbox','notification-service/src/main/java/com/ecommerce/notification_service'),('Browser and gateway','api-gateway/src/main/resources'),('Central service properties','ecommerce-config-repo')]:
    label(name);p(path)
h('Controller coverage register',2)
rows=[]
for file in sorted(ROOT.rglob('*Controller.java')):
    if 'target' not in file.parts and '.repair-backup' not in file.parts:
        text=file.read_text(encoding='utf-8'); count=len(re.findall(r'@(Get|Post|Put|Patch|Delete)Mapping\b',text))
        if count:rows.append([file.name,str(count)])
table(['Controller source','Handler methods'],rows)
p('Compatibility aliases can map more than one URL to one handler. Framework endpoints, static files and outbound provider requests are documented separately. The 63 reference entries cover 64 handler methods because the legacy password reset handler shares one entry with the equivalent authentication handler. Aliases and concrete product or order IDs do not require duplicated entries.')
h('A reproducible demonstration sequence',2)
p('1. Start infrastructure and services with the repository run configuration. Confirm database connectivity, Eureka registrations and service health. Configure placeholder secrets with local test values using the documented environment setup; do not commit real credentials.\n2. Use an existing authorized admin account to create an active product with sufficient stock and optionally SAVE10. There is no public API that grants administrator privileges.\n3. Register a customer, log in, retain both tokens and add an owned delivery address. Browse the public catalog and add two units of the product to the cart.\n4. Submit checkout with a fresh Idempotency-Key. Repeat the exact request with the same key and verify the same order ID. Change a key-bound field and verify a conflict. Add to the cart after the snapshot and verify checkout cleanup preserves the newer cart.\n5. For COD, advance PROCESSING then SHIPPED with carrier and tracking, then DELIVERED with cashCollected=true. For online checkout, use the configured Razorpay test flow and verify the captured payment before fulfillment.\n6. Create a verified-purchase review, request a return, approve it and record physical receipt with an explicit restock decision. Confirm inventory is restored at most once. COD requires an externally paid refund reference; online refunds can remain pending until provider reconciliation.\n7. Read the notification inbox and inspect order event persistence. Repeat selected requests and simulate a temporary dependency outage to explain recovery. Use test data because these operations change orders and stock.')
h('Evidence and maintenance',2)
p('The prior full verification log dated 12 September 2026 records successful Maven verification across all eight modules and 84 tests. That is historical test evidence, not a new execution or a guarantee that a running deployment is healthy on 15 September. Update this guide when DTO fields, mappings, enum values, security matchers, provider validation, migrations or state transitions change. The build inputs in docs/handbook-build preserve the example catalog and narrative for regeneration.')
d.core_properties.title='Ecommerce Application Architecture API Reference and Interview Guide'
d.core_properties.subject='Source based architecture control flow API contracts conditions and interview preparation'
d.core_properties.author='Project documentation'
d.save(OUT)
print(json.dumps({'output':str(OUT),'apis':len(APIS),'questions':len(QUESTIONS),'controller_handlers':sum(int(r[1]) for r in rows),'words':sum(len(p.text.split()) for p in d.paragraphs)}))




